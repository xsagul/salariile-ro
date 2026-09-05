from pathlib import Path
from io import BytesIO
import json,re
from openpyxl import load_workbook
r=Path(__file__).parent
w=load_workbook(BytesIO((r/'raw/uauim-pay-2026.xml').read_bytes()),data_only=True)
out=[]
for s in w:
 rows=list(s.iter_rows(values_only=True))
 print('SHEET',s.title,'rows',s.max_row,'cols',s.max_column)
 for i,row in enumerate(rows):
  if i<14 or re.search('instalator|zugrav|electrician|contabil',str(row),re.I):
   print(i+1,repr(row))
   out.append({'sheet':s.title,'row':i+1,'values':list(row)})
(r/'uauim-extracted-rows.json').write_text(json.dumps(out,ensure_ascii=False,indent=2,default=str),encoding='utf8')
